"""Resume text extraction and structuring.

Pipeline: bytes -> plain text (pdfplumber / python-docx) -> rule-based sections and
entries -> optional LLM refinement whose output is *grounded*: any employer,
title, institution or skill that does not appear in the source text is dropped.
The result is a draft for the user to verify; nothing is written to the profile here.
"""
from __future__ import annotations

import io
import re
from datetime import date

from apps.core import llm
from apps.core.skills import canonicalize, category_of, extract_skills
from apps.core.text import clean_text, normalize_token

# --------------------------------------------------------------------------- text


def extract_text(kind: str, data: bytes) -> str:
    if kind == "pdf":
        import pdfplumber

        pages = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text(x_tolerance=1.5, y_tolerance=3) or "")
        return clean_text("\n".join(pages))
    if kind == "docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    if cell.text not in cells:
                        cells.append(cell.text)
                lines.append(" | ".join(cells))
        return clean_text("\n".join(lines))
    raise ValueError(f"Unsupported resume type: {kind}")


# ----------------------------------------------------------------------- sections

SECTION_ALIASES = {
    "summary": ["summary", "professional summary", "profile", "about me", "career objective", "objective", "career summary"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "work history",
                   "career history", "employment"],
    "education": ["education", "academic background", "academics", "educational qualifications", "qualifications"],
    "skills": ["skills", "technical skills", "core skills", "key skills", "technologies", "tech stack", "core competencies",
               "skills & tools", "technical expertise"],
    "projects": ["projects", "key projects", "personal projects", "selected projects", "academic projects"],
    "certifications": ["certifications", "certificates", "licenses & certifications", "courses", "certifications & training"],
    "achievements": ["achievements", "awards", "honors", "accomplishments", "awards & achievements"],
}
_HEADING_LOOKUP = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}
BULLET_RE = re.compile(r"^\s*(?:[•●▪◦‣∙·\-\*–]|\d{1,2}[.)])\s+")


def _heading_key(line: str):
    candidate = normalize_token(re.sub(r"[:\-–—_#=*|]+\s*$", "", line)).strip(" :")
    if len(candidate) > 40:
        return None
    return _HEADING_LOOKUP.get(candidate)


def split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        key = _heading_key(line) if line else None
        if key:
            current = key
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


# -------------------------------------------------------------------------- dates

_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "sept": 9,
    "oct": 10, "nov": 11, "dec": 12,
}
_MONTH_RE = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
_DATE_RE = rf"(?:{_MONTH_RE}\.?,?\s*'?\d{{2,4}}|\d{{1,2}}[/\-.]\d{{4}}|\d{{4}})"
RANGE_RE = re.compile(
    rf"(?P<start>{_DATE_RE})\s*(?:-|–|—|to|until|till)\s*(?P<end>{_DATE_RE}|present|current|now|till date|today|ongoing)",
    re.I,
)


def parse_date(value: str) -> date | None:
    value = (value or "").strip().lower().replace(".", " ").replace(",", " ")
    if not value:
        return None
    m = re.match(rf"({_MONTH_RE})\s*'?(\d{{2,4}})", value, re.I)
    if m:
        month = _MONTHS[m.group(1)[:4] if m.group(1)[:4] == "sept" else m.group(1)[:3]]
        year = int(m.group(2))
        year = year + 2000 if year < 100 else year
        return date(year, month, 1) if 1950 <= year <= 2100 else None
    m = re.match(r"(\d{1,2})[/\-](\d{4})", value)
    if m and 1 <= int(m.group(1)) <= 12:
        return date(int(m.group(2)), int(m.group(1)), 1)
    m = re.match(r"(\d{4})", value)
    if m and 1950 <= int(m.group(1)) <= 2100:
        return date(int(m.group(1)), 1, 1)
    return None


def _is_present(value: str) -> bool:
    return value.strip().lower() in {"present", "current", "now", "till date", "today", "ongoing"}


def months_between(start: date, end: date) -> int:
    return max(0, (end.year - start.year) * 12 + (end.month - start.month) + 1)


def total_experience_years(entries: list[dict], today: date | None = None) -> float:
    """Merge overlapping employment intervals so concurrent roles are not double counted."""
    today = today or date.today()
    intervals = []
    for e in entries:
        start = date.fromisoformat(e["start_date"]) if e.get("start_date") else None
        if not start:
            continue
        end = today if e.get("is_current") or not e.get("end_date") else date.fromisoformat(e["end_date"])
        if end >= start:
            intervals.append((start, end))
    intervals.sort()
    merged: list[list[date]] = []
    for start, end in intervals:
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    months = sum(months_between(s, e) for s, e in merged)
    return round(months / 12, 1)


# ---------------------------------------------------------------------- entries

TITLE_WORDS = re.compile(
    r"\b(engineer|developer|programmer|architect|lead|manager|consultant|analyst|scientist|intern|trainee|"
    r"specialist|head|director|officer|administrator|designer|sde|swe|associate|member of technical staff|cto|founder)\b",
    re.I,
)
STRONG_COMPANY = re.compile(r"\b(pvt|private|ltd|limited|inc|llc|llp|corp|corporation|gmbh|plc)\b\.?", re.I)
WEAK_COMPANY = re.compile(r"\b(technologies|solutions|systems|labs|software|group|services|consulting|analytics|bank)\b", re.I)
SPLIT_RE = re.compile(r"\s*(?:\||\u2022|·|—|–|\s-\s|,|\bat\b|@)\s*", re.I)


def _strip_bullet(line: str) -> str:
    return BULLET_RE.sub("", line).strip()


def _split_title_company(header: str):
    parts = [p.strip(" ()[]") for p in SPLIT_RE.split(header) if p and p.strip(" ()[]")]
    title = company = location = ""
    leftovers = []
    for part in parts:
        is_title = bool(TITLE_WORDS.search(part)) and not STRONG_COMPANY.search(part)
        is_company = bool(STRONG_COMPANY.search(part)) or (bool(WEAK_COMPANY.search(part)) and not TITLE_WORDS.search(part))
        if is_title and not title:
            title = part
        elif is_company and not company:
            company = part
        elif _looks_like_location(part) and not location:
            location = part
        else:
            leftovers.append(part)
    if not company and leftovers:
        company = leftovers.pop(0)
    if not title and leftovers:
        title = leftovers.pop(0)
    return title[:200], company[:200], location[:200]


_LOCATION_WORDS = re.compile(r"\b(remote|india|noida|delhi|gurgaon|gurugram|bangalore|bengaluru|pune|hyderabad|mumbai|chennai|ncr|usa|uk|singapore|berlin|london)\b", re.I)


def _looks_like_location(part: str) -> bool:
    return bool(_LOCATION_WORDS.search(part)) and len(part.split()) <= 4


def parse_experience(lines: list[str]) -> list[dict]:
    anchors = [i for i, line in enumerate(lines) if line and RANGE_RE.search(line) and not BULLET_RE.match(line)]
    entries = []
    consumed_header_lines: set[int] = set()
    header_ranges = []
    for n, idx in enumerate(anchors):
        line = lines[idx]
        match = RANGE_RE.search(line)
        remainder = (line[: match.start()] + " " + line[match.end():]).strip(" |,-–—()")
        header_lines = [remainder] if remainder else []
        floor = anchors[n - 1] if n else -1
        title, company, location = _split_title_company(remainder) if remainder else ("", "", "")
        j = idx - 1
        taken = 0
        while (not title or not company) and j > floor and taken < 2:
            prev = lines[j]
            if not prev or BULLET_RE.match(prev) or len(prev) > 120:
                break
            header_lines.insert(0, prev)
            consumed_header_lines.add(j)
            taken += 1
            title, company, location = _split_title_company(" | ".join(header_lines))
            j -= 1
        start = parse_date(match.group("start"))
        is_current = _is_present(match.group("end"))
        end = None if is_current else parse_date(match.group("end"))
        if end and start and end < start:
            end = None
        entries.append({
            "company": company, "title": title, "location": location,
            "start_date": start.isoformat() if start else None,
            "end_date": end.isoformat() if end else None,
            "is_current": is_current,
            "description": "", "achievements": [], "technologies": [],
        })
        header_ranges.append(idx)
    for n, idx in enumerate(header_ranges):
        stop = header_ranges[n + 1] if n + 1 < len(header_ranges) else len(lines)
        body = [lines[k] for k in range(idx + 1, stop) if k not in consumed_header_lines and lines[k]]
        achievements = [_strip_bullet(b) for b in body if BULLET_RE.match(b)]
        prose = [b for b in body if not BULLET_RE.match(b)]
        entries[n]["achievements"] = [a for a in achievements if a][:25]
        entries[n]["description"] = " ".join(prose)[:2000]
        entries[n]["technologies"] = extract_skills(" ".join(body))[:30]
    return entries


DEGREE_RE = re.compile(
    r"\b(b\.?\s?tech|m\.?\s?tech|b\.?\s?e\.?|m\.?\s?e\.?|b\.?\s?sc|m\.?\s?sc|bca|mca|mba|ph\.?\s?d|bachelor[’'`s]*|master[’'`s]*|"
    r"diploma|b\.?\s?com|m\.?\s?com|b\.?\s?a\.?|m\.?\s?a\.?|higher secondary|senior secondary|hsc|ssc|12th|10th)\b",
    re.I,
)
INSTITUTION_RE = re.compile(r"\b(university|institute|college|school|iit|nit|iiit|academy|vidyalaya)\b", re.I)
YEAR_RE = re.compile(r"\b(19[5-9]\d|20\d{2})\b")


def parse_education(lines: list[str]) -> list[dict]:
    entries = []
    content = [ln for ln in lines if ln]
    used = set()
    for i, line in enumerate(content):
        if i in used:
            continue
        degree_match = DEGREE_RE.search(line)
        inst_here = INSTITUTION_RE.search(line)
        if not degree_match and not inst_here:
            continue
        window = [line]
        used.add(i)
        if i + 1 < len(content) and (INSTITUTION_RE.search(content[i + 1]) or DEGREE_RE.search(content[i + 1])) \
                and not (degree_match and DEGREE_RE.search(content[i + 1])):
            window.append(content[i + 1])
            used.add(i + 1)
        joined = " | ".join(window)
        years = [int(y) for y in YEAR_RE.findall(joined)]
        institution = degree = field = ""
        for part in [p.strip() for p in re.split(r"\s*(?:\||,|–|—|\s-\s)\s*", joined) if p.strip()]:
            if not institution and INSTITUTION_RE.search(part):
                institution = YEAR_RE.sub("", part).strip(" ()")
            elif not degree and DEGREE_RE.search(part):
                degree = YEAR_RE.sub("", part).strip(" ()")
                m = re.search(r"\b(?:in|of)\s+(.+)$", degree, re.I)
                if m:
                    field = m.group(1).strip()
        grade = ""
        g = re.search(r"\b(cgpa|gpa|percentage)\s*[:\-]?\s*([\d.]+%?(?:\s*/\s*\d+)?)", joined, re.I)
        if g:
            grade = f"{g.group(1).upper()} {g.group(2)}"
        entries.append({
            "institution": institution[:200] or (degree and "") or YEAR_RE.sub("", window[0])[:200],
            "degree": degree[:200], "field_of_study": field[:200],
            "start_year": min(years) if len(years) > 1 else None,
            "end_year": max(years) if years else None,
            "grade": grade[:50],
        })
    return [e for e in entries if e["institution"] or e["degree"]][:10]


def parse_projects(lines: list[str]) -> list[dict]:
    projects: list[dict] = []
    for line in lines:
        if not line:
            continue
        if BULLET_RE.match(line) and projects:
            projects[-1]["highlights"].append(_strip_bullet(line))
        elif not BULLET_RE.match(line) and len(line) <= 120 and (not projects or projects[-1]["highlights"] or projects[-1]["description"]):
            name, _, rest = line.partition(" - ") if " - " in line else line.partition(":")
            projects.append({"name": name.strip()[:200], "role": "", "description": rest.strip(), "highlights": [], "technologies": [], "url": ""})
        elif projects:
            projects[-1]["description"] = (projects[-1]["description"] + " " + _strip_bullet(line)).strip()
        else:
            projects.append({"name": _strip_bullet(line)[:200], "role": "", "description": "", "highlights": [], "technologies": [], "url": ""})
    for p in projects:
        text = " ".join([p["name"], p["description"], *p["highlights"]])
        p["technologies"] = extract_skills(text)[:20]
        url = re.search(r"https?://\S+", text)
        p["url"] = url.group(0).rstrip(").,") if url else ""
    return projects[:20]


def parse_certifications(lines: list[str]) -> list[dict]:
    out = []
    for line in lines:
        text = _strip_bullet(line)
        if not text or len(text) > 200:
            continue
        issuer = ""
        m = re.split(r"\s+(?:-|–|—|by|from|\|)\s+", text, maxsplit=1)
        name = m[0]
        if len(m) > 1:
            issuer = YEAR_RE.sub("", m[1]).strip(" ()")
        out.append({"name": name[:200], "issuer": issuer[:200], "issued_on": None, "credential_url": ""})
    return out[:20]


def parse_skills_section(lines: list[str]) -> list[str]:
    names = []
    for line in lines:
        text = _strip_bullet(line)
        if ":" in text:
            text = text.split(":", 1)[1]
        for token in re.split(r"[,|;/•·]|\s{2,}", text):
            token = token.strip(" .()")
            if 1 < len(token) <= 40:
                names.append(canonicalize(token))
    seen, out = set(), []
    for n in names:
        if n.lower() not in seen:
            seen.add(n.lower())
            out.append(n)
    return out[:150]


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s\-]?)?(?:\(?\d{2,5}\)?[\s\-]?)?\d{3,5}[\s\-]?\d{4,5}")
EXPERIENCE_STATEMENT_RE = re.compile(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years|yrs)\b", re.I)


def parse_contact(header_lines: list[str], text: str) -> dict:
    email = EMAIL_RE.search(text)
    phone = None
    for line in header_lines[:12] or text.splitlines()[:12]:
        m = PHONE_RE.search(line)
        if m and len(re.sub(r"\D", "", m.group(0))) >= 10:
            phone = m.group(0)
            break
    linkedin = re.search(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[A-Za-z0-9\-_%]+/?", text, re.I)
    github = re.search(r"(?:https?://)?github\.com/[A-Za-z0-9\-_]+/?", text, re.I)
    name = ""
    for line in header_lines[:6]:
        words = line.split()
        if 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z.'\-]+", w) for w in words) and not TITLE_WORDS.search(line):
            name = line.title() if line.isupper() else line
            break
    location = ""
    for line in header_lines[:8]:
        m = _LOCATION_WORDS.search(line)
        if m and "@" not in line:
            candidates = [p.strip() for p in re.split(r"[|•·]", line) if _LOCATION_WORDS.search(p)]
            location = candidates[0][:120] if candidates else ""
            break

    def norm_url(m):
        if not m:
            return ""
        url = m.group(0).rstrip("/")
        return url if url.startswith("http") else "https://" + url

    return {
        "full_name": name, "email": email.group(0).lower() if email else "", "phone": phone or "",
        "linkedin_url": norm_url(linkedin), "github_url": norm_url(github), "portfolio_url": "", "location": location,
    }


# --------------------------------------------------------------------- pipeline


def parse_resume_text(text: str, use_llm: bool = True) -> dict:
    sections = split_sections(text)
    header = [ln for ln in sections.get("header", []) if ln]
    contact = parse_contact(header, text)
    employment = parse_experience(sections.get("experience", []))
    summary = " ".join(ln for ln in sections.get("summary", []) if ln)[:3000]
    section_skills = parse_skills_section(sections.get("skills", []))
    mentioned = extract_skills(text)
    skill_names = list(dict.fromkeys(section_skills + mentioned))
    stated = EXPERIENCE_STATEMENT_RE.search(summary or " ".join(header))
    achievements = [_strip_bullet(ln) for ln in sections.get("achievements", []) if ln]
    headline = ""
    for line in header[1:5]:
        if TITLE_WORDS.search(line) and "@" not in line and len(line) < 120:
            headline = line
            break

    result = {
        "contact": contact,
        "headline": headline,
        "summary": summary,
        "skills": [
            {"name": n, "category": category_of(n), "in_skills_section": n in section_skills} for n in skill_names
        ],
        "employment": employment,
        "education": parse_education(sections.get("education", [])),
        "projects": parse_projects(sections.get("projects", [])),
        "certifications": parse_certifications(sections.get("certifications", [])),
        "achievements": achievements[:20],
        "total_experience_years": total_experience_years(employment),
        "stated_experience_years": float(stated.group(1)) if stated else None,
        "warnings": [],
        "parser": "rules",
    }
    if not employment:
        result["warnings"].append("No employment entries with date ranges were detected; please add them manually.")
    if use_llm and llm.is_enabled():
        result = _refine_with_llm(text, result)
    return result


# ---------------------------------------------------------------- LLM refinement

_STR = {"type": "string"}
_NULLABLE_STR = {"type": ["string", "null"]}
LLM_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["headline", "summary", "employment", "education", "skills"],
    "properties": {
        "headline": _STR,
        "summary": _STR,
        "skills": {"type": "array", "items": _STR},
        "employment": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["company", "title", "location", "start_date", "end_date", "is_current", "achievements"],
                "properties": {
                    "company": _STR, "title": _STR, "location": _STR,
                    "start_date": _NULLABLE_STR, "end_date": _NULLABLE_STR, "is_current": {"type": "boolean"},
                    "achievements": {"type": "array", "items": _STR},
                },
            },
        },
        "education": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["institution", "degree", "field_of_study", "end_year"],
                "properties": {"institution": _STR, "degree": _STR, "field_of_study": _STR, "end_year": {"type": ["integer", "null"]}},
            },
        },
    },
}


def _grounded(value: str, text_norm: str) -> bool:
    v = normalize_token(re.sub(r"[^\w\s+#.]", " ", value or ""))
    return bool(v) and v in text_norm


def _refine_with_llm(text: str, rules: dict) -> dict:
    instructions = (
        "Extract the candidate's resume into the JSON schema. Copy employer names, job titles, institutions and "
        "skills exactly as written in the resume. Dates as YYYY-MM-01 (or null if absent). Do not infer, embellish "
        "or add anything that is not explicitly stated in the resume."
    )
    try:
        result = llm.generate_json("resume_extraction", instructions, LLM_SCHEMA, untrusted_blocks={"resume": text})
    except llm.LLMUnavailable:
        rules["warnings"].append("AI extraction unavailable; used rule-based parsing only.")
        return rules
    data = result.data
    text_norm = normalize_token(re.sub(r"[^\w\s+#.]", " ", text))
    dropped = []
    employment = []
    for e in data.get("employment", []):
        if not (_grounded(e.get("company", ""), text_norm) and _grounded(e.get("title", ""), text_norm)):
            dropped.append(f"employment entry '{e.get('title')} @ {e.get('company')}'")
            continue
        for key in ("start_date", "end_date"):
            try:
                e[key] = date.fromisoformat(e[key]).replace(day=1).isoformat() if e.get(key) else None
            except ValueError:
                e[key] = None
        e["achievements"] = [a for a in e.get("achievements", []) if a][:25]
        e["description"] = ""
        e["technologies"] = extract_skills(" ".join(e["achievements"]))[:30]
        employment.append(e)
    education = []
    for ed in data.get("education", []):
        if _grounded(ed.get("institution", ""), text_norm):
            education.append({**ed, "start_year": None, "grade": ""})
        else:
            dropped.append(f"education entry '{ed.get('institution')}'")
    llm_skills = [canonicalize(s) for s in data.get("skills", []) if _grounded(s, text_norm)]

    merged = dict(rules)
    if employment and len(employment) >= len(rules["employment"]):
        merged["employment"] = employment
        merged["total_experience_years"] = total_experience_years(employment)
    if education and len(education) >= len(rules["education"]):
        merged["education"] = education
    known = {s["name"] for s in rules["skills"]}
    merged["skills"] = rules["skills"] + [
        {"name": s, "category": category_of(s), "in_skills_section": False} for s in dict.fromkeys(llm_skills) if s not in known
    ]
    if data.get("headline") and _grounded(data["headline"], text_norm):
        merged["headline"] = data["headline"][:200]
    if dropped:
        merged["warnings"] = rules["warnings"] + [f"Ignored AI output not found in the resume: {', '.join(dropped[:5])}"]
    merged["parser"] = f"rules+{result.provider}"
    return merged
