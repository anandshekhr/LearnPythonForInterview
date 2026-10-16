"""Render structured resume content / cover letters to DOCX and PDF using approved templates."""
from __future__ import annotations

import io
from datetime import date
from xml.sax.saxutils import escape

TEMPLATES = {
    "classic": {"font": "Helvetica", "bold": "Helvetica-Bold", "docx_font": "Calibri", "base": 10, "heading": 12, "name": 18,
                "accent": (0.12, 0.25, 0.45)},
    "compact": {"font": "Times-Roman", "bold": "Times-Bold", "docx_font": "Georgia", "base": 9.5, "heading": 11, "name": 16,
                "accent": (0.1, 0.1, 0.1)},
}


def _period(e):
    def fmt(value):
        if not value:
            return ""
        d = date.fromisoformat(value)
        return d.strftime("%b %Y")
    end = "Present" if e.get("is_current") else fmt(e.get("end_date"))
    start = fmt(e.get("start_date"))
    return f"{start} – {end}" if start or end else ""


def _contact_line(c):
    return " | ".join(x for x in [c.get("email"), c.get("phone"), c.get("location"), c.get("linkedin_url"), c.get("github_url"),
                                  c.get("portfolio_url")] if x)


def resume_docx(content: dict, template="classic") -> bytes:
    import docx
    from docx.shared import Pt, RGBColor

    t = TEMPLATES[template]
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = t["docx_font"]
    style.font.size = Pt(t["base"])
    for section in document.sections:
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(42)
    c = content.get("contact", {})
    name = document.add_paragraph()
    run = name.add_run(c.get("full_name") or "")
    run.bold, run.font.size = True, Pt(t["name"])
    if content.get("headline"):
        document.add_paragraph(content["headline"])
    document.add_paragraph(_contact_line(c))

    def heading(text):
        p = document.add_paragraph()
        r = p.add_run(text.upper())
        r.bold, r.font.size = True, Pt(t["heading"])
        r.font.color.rgb = RGBColor(*[int(x * 255) for x in t["accent"]])

    if content.get("summary"):
        heading("Summary")
        document.add_paragraph(content["summary"])
    if content.get("skills"):
        heading("Skills")
        document.add_paragraph(", ".join(content["skills"]))
    if content.get("employment"):
        heading("Experience")
        for e in content["employment"]:
            p = document.add_paragraph()
            r = p.add_run(f"{e['title']} — {e['company']}")
            r.bold = True
            meta = " · ".join(x for x in [e.get("location"), _period(e)] if x)
            if meta:
                p.add_run(f"   {meta}")
            if e.get("description"):
                document.add_paragraph(e["description"])
            for a in e.get("achievements", []):
                document.add_paragraph(a, style="List Bullet")
    if content.get("projects"):
        heading("Projects")
        for pr in content["projects"]:
            p = document.add_paragraph()
            p.add_run(pr["name"]).bold = True
            if pr.get("description"):
                p.add_run(f" — {pr['description']}")
            for h in pr.get("highlights", []):
                document.add_paragraph(h, style="List Bullet")
    if content.get("education"):
        heading("Education")
        for ed in content["education"]:
            line = ", ".join(x for x in [ed.get("degree"), ed.get("institution"), str(ed["end_year"]) if ed.get("end_year") else "",
                                         ed.get("grade")] if x)
            document.add_paragraph(line)
    if content.get("certifications"):
        heading("Certifications")
        for cert in content["certifications"]:
            document.add_paragraph(" – ".join(x for x in [cert["name"], cert.get("issuer")] if x), style="List Bullet")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _pdf(story_builder, template="classic") -> bytes:
    from reportlab.lib.colors import Color
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate

    t = TEMPLATES[template]
    styles = {
        "body": ParagraphStyle("body", fontName=t["font"], fontSize=t["base"], leading=t["base"] * 1.35),
        "name": ParagraphStyle("name", fontName=t["bold"], fontSize=t["name"], leading=t["name"] * 1.2),
        "heading": ParagraphStyle("heading", fontName=t["bold"], fontSize=t["heading"], leading=t["heading"] * 1.6,
                                  textColor=Color(*t["accent"]), spaceBefore=6),
        "bullet": ParagraphStyle("bullet", fontName=t["font"], fontSize=t["base"], leading=t["base"] * 1.35, leftIndent=10,
                                 bulletIndent=2),
    }
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title="Resume", author="")
    doc.build(story_builder(styles))
    return buf.getvalue()


def resume_pdf(content: dict, template="classic") -> bytes:
    from reportlab.platypus import Paragraph, Spacer

    def build(st):
        c = content.get("contact", {})
        story = [Paragraph(escape(c.get("full_name") or ""), st["name"])]
        if content.get("headline"):
            story.append(Paragraph(escape(content["headline"]), st["body"]))
        story.append(Paragraph(escape(_contact_line(c)), st["body"]))
        if content.get("summary"):
            story += [Paragraph("SUMMARY", st["heading"]), Paragraph(escape(content["summary"]), st["body"])]
        if content.get("skills"):
            story += [Paragraph("SKILLS", st["heading"]), Paragraph(escape(", ".join(content["skills"])), st["body"])]
        if content.get("employment"):
            story.append(Paragraph("EXPERIENCE", st["heading"]))
            for e in content["employment"]:
                meta = " · ".join(x for x in [e.get("location"), _period(e)] if x)
                story.append(Paragraph(f"<b>{escape(e['title'])} — {escape(e['company'])}</b>  {escape(meta)}", st["body"]))
                if e.get("description"):
                    story.append(Paragraph(escape(e["description"]), st["body"]))
                for a in e.get("achievements", []):
                    story.append(Paragraph(escape(a), st["bullet"], bulletText="•"))
                story.append(Spacer(1, 4))
        if content.get("projects"):
            story.append(Paragraph("PROJECTS", st["heading"]))
            for p in content["projects"]:
                story.append(Paragraph(f"<b>{escape(p['name'])}</b> {escape(p.get('description') or '')}", st["body"]))
                for h in p.get("highlights", []):
                    story.append(Paragraph(escape(h), st["bullet"], bulletText="•"))
        if content.get("education"):
            story.append(Paragraph("EDUCATION", st["heading"]))
            for ed in content["education"]:
                line = ", ".join(x for x in [ed.get("degree"), ed.get("institution"),
                                             str(ed["end_year"]) if ed.get("end_year") else "", ed.get("grade")] if x)
                story.append(Paragraph(escape(line), st["body"]))
        if content.get("certifications"):
            story.append(Paragraph("CERTIFICATIONS", st["heading"]))
            for cert in content["certifications"]:
                story.append(Paragraph(escape(" – ".join(x for x in [cert["name"], cert.get("issuer")] if x)), st["bullet"],
                                       bulletText="•"))
        return story

    return _pdf(build, template)


def letter_docx(text: str, template="classic") -> bytes:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    document.styles["Normal"].font.name = TEMPLATES[template]["docx_font"]
    document.styles["Normal"].font.size = Pt(11)
    for para in text.split("\n\n"):
        document.add_paragraph(para.strip())
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def letter_pdf(text: str, template="classic") -> bytes:
    from reportlab.platypus import Paragraph, Spacer

    def build(st):
        story = []
        for para in text.split("\n\n"):
            story += [Paragraph(escape(para.strip()).replace("\n", "<br/>"), st["body"]), Spacer(1, 8)]
        return story

    return _pdf(build, template)
